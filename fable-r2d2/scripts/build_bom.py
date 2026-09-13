"""Build bom/bill-of-materials.xlsx from the CSV sources with live totals and source links.

Inputs, all read at build time: bom/electronics.csv, bom/hardware.csv (columns: reference,
quantity, item, part_number, supplier, url, unit_usd, price_basis, notes), bom/printed-parts.csv
(written by scripts/export_cad.py), cad/h2d-slice-check.json for filament mass, and the frozen
previous revision named in docs/revision-c.json (read from git) for the revision D change list.
Output: bom/bill-of-materials.xlsx with sheets Summary, Electronics, Hardware, Printed parts,
Filament, Revision D changes, Stock notes, Sources. Formulas compute line and sheet totals so the
operator can edit quantities or prices in Excel and see totals update.
"""
import csv
import io
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from assembly_layout import REVISION, read_parts  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "bom/bill-of-materials.xlsx"
BASELINE = ROOT / "docs/revision-c.json"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="B8C2CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COLUMNS = ["reference", "quantity", "item", "part_number", "supplier", "url", "unit_usd", "price_basis", "notes"]
TITLES = ["Ref", "Qty", "Item", "Part number", "Supplier", "Source URL", "Unit USD", "Price basis", "Notes / use"]
STOCK_RE = re.compile(r"[^.;]*\b(out of stock|back order|backorder|no longer stocked|discontinued)\b[^.;]*", re.I)


def read_csv(name, root=ROOT):
    path = Path(root) / "bom" / name
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def baseline(root=ROOT):
    """The frozen previous revision: {"revision": "C", "git_tag": ...} from docs/revision-c.json."""
    data = json.loads((Path(root) / "docs/revision-c.json").read_text(encoding="utf-8"))
    return {"revision": data["revision"], "git_tag": data["git_tag"]}


def git_show(relative, root=ROOT):
    """A package file's text at the baseline tag."""
    root = Path(root)
    prefix = subprocess.check_output(["git", "rev-parse", "--show-prefix"], cwd=root, text=True).strip()
    tag = baseline(root)["git_tag"]
    return subprocess.check_output(["git", "show", f"{tag}:{prefix}{relative}"], cwd=root, text=True, encoding="utf-8")


def _number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def purchased_delta(name, root=ROOT):
    """Rows of bom/<name> added, changed or removed since the baseline revision.

    Rows are matched by reference; a new reference with no old match is matched to a removed old
    row with the same part number (a renamed designator such as FH1-FH6 -> FH1-FH7).
    """
    old = list(csv.DictReader(io.StringIO(git_show(f"bom/{name}", root))))
    new = read_csv(name, root)
    old_by_ref = {r["reference"]: r for r in old}
    new_refs = {r["reference"] for r in new}
    removed = {r["reference"]: r for r in old if r["reference"] not in new_refs}
    rows = []
    for row in new:
        previous = old_by_ref.get(row["reference"])
        if previous is None:
            previous = next((r for r in removed.values() if r.get("part_number") == row.get("part_number")), None)
            if previous is not None:
                removed.pop(previous["reference"])
        old_qty = _number(previous["quantity"]) if previous else 0.0
        new_qty = _number(row["quantity"])
        if previous is not None and old_qty == new_qty and previous["reference"] == row["reference"]:
            continue
        if previous is not None and old_qty == new_qty:
            status = "renamed"
        else:
            status = "added" if previous is None else "changed"
        rows.append({"reference": row["reference"], "previous_reference": previous["reference"] if previous else "",
                     "item": row["item"], "part_number": row.get("part_number", ""), "old_quantity": old_qty,
                     "new_quantity": new_qty, "delta_quantity": new_qty - old_qty, "unit_usd": _number(row["unit_usd"]),
                     "status": status})
    for ref, row in removed.items():
        rows.append({"reference": "", "previous_reference": ref, "item": row["item"],
                     "part_number": row.get("part_number", ""), "old_quantity": _number(row["quantity"]),
                     "new_quantity": 0.0, "delta_quantity": -_number(row["quantity"]),
                     "unit_usd": _number(row["unit_usd"]), "status": "removed"})
    for row in rows:
        row["delta_usd"] = round(row["delta_quantity"] * row["unit_usd"], 2)
    return {"file": f"bom/{name}", "rows": rows,
            "pieces": sum(r["delta_quantity"] for r in rows),
            "usd": round(sum(r["delta_usd"] for r in rows), 2)}


def printed_delta(root=ROOT):
    """Printed designs and pieces added, changed or removed since the baseline revision."""
    old = json.loads(git_show("scripts/parts.json", root))["parts"]
    new = read_parts(root)
    rows = []
    for name in sorted(set(old) | set(new)):
        before = int(old[name]["quantity"]) if name in old else 0
        after = int(new[name]["quantity"]) if name in new else 0
        if name in old and name in new and before == after and old[name].get("notes") == new[name].get("notes"):
            continue
        status = "added" if name not in old else "removed" if name not in new else "changed"
        rows.append({"part": name, "status": status, "old_quantity": before, "new_quantity": after,
                     "notes": (new.get(name) or old.get(name)).get("notes", "")})
    return {"rows": rows, "designs": {"before": len(old), "after": len(new)},
            "pieces": {"before": sum(int(r["quantity"]) for r in old.values()),
                       "after": sum(int(r["quantity"]) for r in new.values())}}


def stock_notes(root=ROOT):
    """Purchased rows whose notes record a stock problem, with the sentence that says so."""
    found = []
    for name in ("electronics.csv", "hardware.csv"):
        for row in read_csv(name, root):
            match = STOCK_RE.search(row.get("notes", ""))
            if match:
                found.append({"file": f"bom/{name}", "reference": row["reference"], "item": row["item"],
                              "part_number": row.get("part_number", ""), "supplier": row.get("supplier", ""),
                              "url": row.get("url", ""), "status": match.group(1).lower(),
                              "note": match.group(0).strip()})
    return found


def style_header(sheet, titles):
    for col, title in enumerate(titles, 1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 30


def purchased_sheet(book, title, rows):
    sheet = book.create_sheet(title)
    style_header(sheet, TITLES + ["Line USD"])
    for index, row in enumerate(rows, 2):
        for col, key in enumerate(COLUMNS, 1):
            value = row.get(key, "")
            if key == "quantity":
                value = int(float(value)) if value else 0
            elif key == "unit_usd":
                value = float(value) if value else 0.0
            cell = sheet.cell(row=index, column=col, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("item", "notes"))
            if key == "url" and value:
                cell.hyperlink = value
                cell.font = Font(color="0B5394", underline="single")
            if key == "unit_usd":
                cell.number_format = "$#,##0.00"
        line = sheet.cell(row=index, column=len(COLUMNS) + 1, value=f"=B{index}*G{index}")
        line.number_format = "$#,##0.00"
        line.border = BORDER
    total_row = len(rows) + 2
    sheet.cell(row=total_row, column=3, value="Sheet total").font = TOTAL_FONT
    total = sheet.cell(row=total_row, column=len(COLUMNS) + 1, value=f"=SUM(J2:J{total_row - 1})")
    total.number_format = "$#,##0.00"
    total.font = TOTAL_FONT
    widths = [8, 6, 34, 20, 14, 46, 10, 12, 52, 11]
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    return sheet, total_row


def printed_sheet(book, rows, slices):
    sheet = book.create_sheet("Printed parts")
    titles = ["STL file", "Qty", "Material", "Mirror for 2nd", "X mm", "Y mm", "Z mm", "Solid mass g (bound)",
              "Sliced mass g (H2D, incl. support)", "Sliced time h", "Orientation", "Notes"]
    style_header(sheet, titles)
    by_part = {Path(r["part"]).stem: r for r in slices.get("rows", [])}
    for index, row in enumerate(rows, 2):
        sliced = by_part.get(row["part"], {})
        values = [f"stl/{row['part']}.stl", int(row["quantity"]), row["material"],
                  "yes" if str(row.get("mirror_for_second", "")).lower() == "true" else "no",
                  float(row["x_mm"]), float(row["y_mm"]), float(row["z_mm"]), float(row["solid_mass_g"]),
                  round(float(sliced.get("predicted_mass_g", 0)), 1), float(sliced.get("predicted_time_h", 0)),
                  row.get("orientation", ""), row.get("notes", "")]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row=index, column=col, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col >= 11)
    total_row = len(rows) + 2
    sheet.cell(row=total_row, column=1, value="Totals (qty x per-part)").font = TOTAL_FONT
    sheet.cell(row=total_row, column=9, value=f"=SUMPRODUCT(B2:B{total_row - 1},I2:I{total_row - 1})").font = TOTAL_FONT
    sheet.cell(row=total_row, column=10, value=f"=SUMPRODUCT(B2:B{total_row - 1},J2:J{total_row - 1})").font = TOTAL_FONT
    for col, width in enumerate([24, 6, 10, 12, 9, 9, 9, 14, 16, 12, 30, 60], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    return sheet, total_row


def filament_sheet(book, rows, slices):
    sheet = book.create_sheet("Filament")
    style_header(sheet, ["Material", "Sliced mass g (qty applied)", "Spools of 1 kg (with 25 % waste)", "Spool USD", "Cost USD", "Source"])
    by_part = {Path(r["part"]).stem: r for r in slices.get("rows", [])}
    materials = {}
    for row in rows:
        mass = float(by_part.get(row["part"], {}).get("predicted_mass_g", 0)) * int(row["quantity"])
        materials[row["material"]] = materials.get(row["material"], 0) + mass
    prices = {"PETG": (24.99, "https://us.store.bambulab.com/products/petg-hf"),
              "PETG-CF": (34.99, "https://us.store.bambulab.com/products/petg-cf"),
              "PLA": (19.99, "https://us.store.bambulab.com/products/pla-basic")}
    for index, (material, mass) in enumerate(sorted(materials.items()), 2):
        price, url = prices.get(material, (25.0, ""))
        sheet.cell(row=index, column=1, value=material)
        sheet.cell(row=index, column=2, value=round(mass, 1))
        sheet.cell(row=index, column=3, value=f"=ROUNDUP(B{index}*1.25/1000,0)")
        sheet.cell(row=index, column=4, value=price).number_format = "$#,##0.00"
        sheet.cell(row=index, column=5, value=f"=C{index}*D{index}").number_format = "$#,##0.00"
        link = sheet.cell(row=index, column=6, value=url)
        if url:
            link.hyperlink = url
            link.font = Font(color="0B5394", underline="single")
    total_row = len(materials) + 2
    sheet.cell(row=total_row, column=1, value="Total").font = TOTAL_FONT
    sheet.cell(row=total_row, column=5, value=f"=SUM(E2:E{total_row - 1})").number_format = "$#,##0.00"
    for col, width in enumerate([12, 24, 30, 12, 12, 50], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    return sheet, total_row


def changes_sheet(book, deltas, printed, base):
    """Purchased and printed changes since the baseline revision, with live line formulas."""
    sheet = book.create_sheet(f"Revision {REVISION} changes")
    style_header(sheet, ["BOM file", "Status", "Ref", f"Revision {base['revision']} ref", "Item", "Part number",
                         f"Qty rev {base['revision']}", f"Qty rev {REVISION}", "Qty change", "Unit USD", "Change USD"])
    row = 2
    first = row
    for delta in deltas:
        for entry in delta["rows"]:
            values = [delta["file"], entry["status"], entry["reference"], entry["previous_reference"], entry["item"],
                      entry["part_number"], entry["old_quantity"], entry["new_quantity"]]
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row=row, column=col, value=value)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=col == 5)
            sheet.cell(row=row, column=9, value=f"=H{row}-G{row}").border = BORDER
            sheet.cell(row=row, column=10, value=entry["unit_usd"]).number_format = "$#,##0.00"
            line = sheet.cell(row=row, column=11, value=f"=I{row}*J{row}")
            line.number_format = "$#,##0.00"
            line.border = BORDER
            row += 1
    total_row = row
    sheet.cell(row=total_row, column=5, value="Purchased change total").font = TOTAL_FONT
    sheet.cell(row=total_row, column=9, value=f"=SUM(I{first}:I{total_row - 1})").font = TOTAL_FONT
    usd = sheet.cell(row=total_row, column=11, value=f"=SUM(K{first}:K{total_row - 1})")
    usd.number_format = "$#,##0.00"
    usd.font = TOTAL_FONT
    row = total_row + 2
    sheet.cell(row=row, column=1, value="Printed parts").font = TOTAL_FONT
    for entry in printed["rows"]:
        row += 1
        for col, value in enumerate(["scripts/parts.json", entry["status"], entry["part"], "", entry["notes"], "",
                                     entry["old_quantity"], entry["new_quantity"]], 1):
            cell = sheet.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=col == 5)
        sheet.cell(row=row, column=9, value=f"=H{row}-G{row}")
    for col, width in enumerate([22, 10, 12, 14, 48, 22, 10, 10, 10, 10, 12], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    return sheet, total_row


def stock_sheet(book, notes):
    sheet = book.create_sheet("Stock notes")
    style_header(sheet, ["BOM file", "Ref", "Item", "Part number", "Supplier", "Status", "What the supplier page read", "Source URL"])
    for index, note in enumerate(notes, 2):
        values = [note["file"], note["reference"], note["item"], note["part_number"], note["supplier"], note["status"],
                  note["note"], note["url"]]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row=index, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (3, 7))
            if col == 8 and value:
                cell.hyperlink = value
                cell.font = Font(color="0B5394", underline="single")
    for col, width in enumerate([18, 10, 40, 22, 14, 16, 60, 40], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width


def sources_sheet(book, electronics, hardware):
    sheet = book.create_sheet("Sources")
    style_header(sheet, ["Supplier", "Items", "URLs"])
    suppliers = {}
    for row in electronics + hardware:
        entry = suppliers.setdefault(row["supplier"], {"items": [], "urls": set()})
        entry["items"].append(row["item"])
        if row.get("url"):
            entry["urls"].add(row["url"])
    for index, (supplier, entry) in enumerate(sorted(suppliers.items()), 2):
        sheet.cell(row=index, column=1, value=supplier)
        sheet.cell(row=index, column=2, value="; ".join(entry["items"])).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row=index, column=3, value="\n".join(sorted(entry["urls"]))).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in enumerate([22, 70, 70], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width


def main():
    electronics = read_csv("electronics.csv")
    hardware = read_csv("hardware.csv")
    printed = read_csv("printed-parts.csv")
    slices = json.loads((ROOT / "cad/h2d-slice-check.json").read_text(encoding="utf-8")) if (ROOT / "cad/h2d-slice-check.json").exists() else {}
    base = baseline()
    deltas = [purchased_delta("electronics.csv"), purchased_delta("hardware.csv")]
    printed_changes = printed_delta()
    notes = stock_notes()
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    _, e_total = purchased_sheet(book, "Electronics", electronics)
    _, h_total = purchased_sheet(book, "Hardware", hardware)
    printed_sheet(book, printed, slices)
    _, f_total = filament_sheet(book, printed, slices)
    change_title = f"Revision {REVISION} changes"
    _, c_total = changes_sheet(book, deltas, printed_changes, base)
    stock_sheet(book, notes)
    sources_sheet(book, electronics, hardware)
    summary["A1"] = f"fable-r2d2 revision {REVISION} bill of materials"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = ("Prices are list prices read from supplier pages on the date in the notes; tax and shipping excluded. "
                     "Edit quantities or unit prices on the sheets; totals are formulas.")
    summary["A2"].alignment = Alignment(wrap_text=True)
    summary.merge_cells("A2:C2")
    summary.row_dimensions[2].height = 45
    lines = [("Electronics", f"=Electronics!J{e_total}"), ("Mechanical hardware", f"=Hardware!J{h_total}"),
             ("Filament", f"=Filament!E{f_total}"), ("Grand total (USD)", "=SUM(B4:B6)")]
    for offset, (label, formula) in enumerate(lines):
        row = 4 + offset
        summary.cell(row=row, column=1, value=label).font = TOTAL_FONT if offset == 3 else Font()
        cell = summary.cell(row=row, column=2, value=formula)
        cell.number_format = "$#,##0.00"
        if offset == 3:
            cell.font = TOTAL_FONT
    summary["A9"] = "Printed pieces"
    summary["B9"] = f"=SUM('Printed parts'!B2:B{len(printed) + 1})"
    summary["A10"] = "Unique STL files"
    summary["B10"] = len(printed)
    summary["A11"] = "Sliced filament mass, g"
    summary["B11"] = f"='Printed parts'!I{len(printed) + 2}"
    summary["A12"] = "Sliced print time, h"
    summary["B12"] = f"='Printed parts'!J{len(printed) + 2}"
    summary["A14"] = f"Purchased pieces added since revision {base['revision']} ({base['git_tag']})"
    summary["B14"] = f"='{change_title}'!I{c_total}"
    summary["A15"] = f"Purchased cost added since revision {base['revision']}, USD"
    summary["B15"] = f"='{change_title}'!K{c_total}"
    summary["B15"].number_format = "$#,##0.00"
    summary["A16"] = "Purchased rows with a stock note"
    summary["B16"] = len(notes)
    summary["C16"] = "; ".join(f"{n['reference']} {n['status']}" for n in notes)
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 60
    OUT.parent.mkdir(exist_ok=True)
    book.save(OUT)
    parts = read_parts()
    lagging = sorted(set(parts) - {row["part"] for row in printed})
    if lagging:
        # export_cad.py writes printed-parts.csv; verify.py check_cad fails the package until it has run.
        print(f"NOTE: bom/printed-parts.csv has no row for {', '.join(lagging)} listed in scripts/parts.json; "
              "the Printed parts and Filament sheets omit it until python scripts/export_cad.py runs")
    print(f"PASS: {OUT} revision {REVISION}: {len(electronics)} electronics rows, {len(hardware)} hardware rows, "
          f"{len(printed)} printed parts; since revision {base['revision']}: electronics "
          f"{deltas[0]['pieces']:+g} pieces ${deltas[0]['usd']:+,.2f}, hardware {deltas[1]['pieces']:+g} pieces "
          f"${deltas[1]['usd']:+,.2f}, printed designs {printed_changes['designs']['before']} -> "
          f"{printed_changes['designs']['after']}; {len(notes)} stock notes")


if __name__ == "__main__":
    main()
